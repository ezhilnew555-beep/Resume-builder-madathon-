"""
Resume Export Service for CareerLens (Phase 9).
Compiles structured resume builder data into authentic, ATS-compliant
PDF and DOCX documents with selectable text, styling, and page-overflow handling.
"""

import io
from typing import Optional
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.builder import ResumeBuilderData

class ResumeExportService:
    def export_docx(self, data: ResumeBuilderData, template_id: str = "modern") -> bytes:
        """
        Compiles ResumeBuilderData into a formatted, ATS-compliant Microsoft Word (.docx) document.
        """
        doc = Document()
        
        # Set page margins to standard 0.75 in
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Template font preferences
        font_name = "Calibri" if template_id in ["modern", "minimal"] else "Times New Roman"
        
        # 1. Header (Name & Contact)
        p_name = doc.add_paragraph()
        p_name.paragraph_format.space_after = Pt(2)
        p_name.paragraph_format.space_before = Pt(0)
        if template_id == "classic":
            p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run_name = p_name.add_run(data.contact.full_name or "Candidate Name")
        run_name.font.name = font_name
        run_name.font.size = Pt(20)
        run_name.font.bold = True
        if template_id == "modern":
            run_name.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

        # Contact line
        contact_parts = []
        if data.contact.location: contact_parts.append(data.contact.location)
        if data.contact.phone: contact_parts.append(data.contact.phone)
        if data.contact.email: contact_parts.append(data.contact.email)
        if data.contact.linkedin: contact_parts.append(data.contact.linkedin.replace("https://", ""))
        if data.contact.github: contact_parts.append(data.contact.github.replace("https://", ""))
        if data.contact.portfolio: contact_parts.append(data.contact.portfolio.replace("https://", ""))

        if contact_parts:
            p_contact = doc.add_paragraph()
            p_contact.paragraph_format.space_after = Pt(8)
            p_contact.paragraph_format.space_before = Pt(0)
            if template_id == "classic":
                p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            run_contact = p_contact.add_run(" • ".join(contact_parts))
            run_contact.font.name = font_name
            run_contact.font.size = Pt(9.5)
            run_contact.font.color.rgb = RGBColor(100, 116, 139)

        # Helper for adding section headings
        def add_heading(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(title.upper())
            run.font.name = font_name
            run.font.size = Pt(11)
            run.font.bold = True
            if template_id == "modern":
                run.font.color.rgb = RGBColor(30, 58, 138)
            else:
                run.font.color.rgb = RGBColor(15, 23, 42)

        # Render sections according to section_order
        for section_key in data.section_order:
            if section_key == "summary" and data.summary.strip():
                add_heading("Professional Summary")
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(data.summary.strip())
                run.font.name = font_name
                run.font.size = Pt(10)

            elif section_key == "skills":
                s = data.skills
                has_skills = any([s.languages, s.frameworks, s.databases, s.tools, s.soft_skills])
                if has_skills:
                    add_heading("Technical Skills")
                    skill_rows = []
                    if s.languages: skill_rows.append(("Languages", s.languages))
                    if s.frameworks: skill_rows.append(("Frameworks & Libraries", s.frameworks))
                    if s.databases: skill_rows.append(("Databases & Cloud", s.databases))
                    if s.tools: skill_rows.append(("Developer Tools", s.tools))
                    if s.soft_skills: skill_rows.append(("Core Competencies", s.soft_skills))

                    for label, val in skill_rows:
                        p = doc.add_paragraph()
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(2)
                        
                        r_lbl = p.add_run(f"{label}: ")
                        r_lbl.font.name = font_name
                        r_lbl.font.size = Pt(9.5)
                        r_lbl.font.bold = True
                        
                        r_val = p.add_run(val)
                        r_val.font.name = font_name
                        r_val.font.size = Pt(9.5)

            elif section_key == "experience" and data.experience:
                add_heading("Work & Internship Experience")
                for exp in data.experience:
                    p_title = doc.add_paragraph()
                    p_title.paragraph_format.space_before = Pt(4)
                    p_title.paragraph_format.space_after = Pt(1)
                    p_title.paragraph_format.keep_with_next = True
                    
                    r_role = p_title.add_run(f"{exp.role} — ")
                    r_role.font.bold = True
                    r_role.font.size = Pt(10.5)
                    r_role.font.name = font_name
                    
                    r_comp = p_title.add_run(exp.company)
                    r_comp.font.size = Pt(10)
                    r_comp.font.name = font_name
                    
                    date_loc = []
                    if exp.start_date or exp.end_date:
                        date_loc.append(f"{exp.start_date} - {exp.end_date if not exp.is_current else 'Present'}")
                    if exp.location:
                        date_loc.append(exp.location)
                    if date_loc:
                        r_date = p_title.add_run(f" ({' | '.join(date_loc)})")
                        r_date.font.size = Pt(9)
                        r_date.font.italic = True
                        r_date.font.name = font_name

                    for bullet in exp.bullets:
                        if bullet.strip():
                            bp = doc.add_paragraph(style='List Bullet')
                            bp.paragraph_format.space_before = Pt(0)
                            bp.paragraph_format.space_after = Pt(2)
                            r_b = bp.add_run(bullet.strip())
                            r_b.font.name = font_name
                            r_b.font.size = Pt(9.5)

            elif section_key == "projects" and data.projects:
                add_heading("Key Technical Projects")
                for proj in data.projects:
                    p_pname = doc.add_paragraph()
                    p_pname.paragraph_format.space_before = Pt(4)
                    p_pname.paragraph_format.space_after = Pt(1)
                    p_pname.paragraph_format.keep_with_next = True
                    
                    r_pname = p_pname.add_run(proj.name)
                    r_pname.font.bold = True
                    r_pname.font.size = Pt(10.5)
                    r_pname.font.name = font_name

                    if proj.tech_stack:
                        r_tech = p_pname.add_run(f" | {proj.tech_stack}")
                        r_tech.font.size = Pt(9.5)
                        r_tech.font.italic = True
                        r_tech.font.name = font_name

                    for bullet in proj.bullets:
                        if bullet.strip():
                            bp = doc.add_paragraph(style='List Bullet')
                            bp.paragraph_format.space_before = Pt(0)
                            bp.paragraph_format.space_after = Pt(2)
                            r_b = bp.add_run(bullet.strip())
                            r_b.font.name = font_name
                            r_b.font.size = Pt(9.5)

            elif section_key == "education" and data.education:
                add_heading("Education")
                for edu in data.education:
                    p_edu = doc.add_paragraph()
                    p_edu.paragraph_format.space_before = Pt(3)
                    p_edu.paragraph_format.space_after = Pt(1)
                    
                    r_deg = p_edu.add_run(f"{edu.degree}")
                    r_deg.font.bold = True
                    r_deg.font.size = Pt(10)
                    r_deg.font.name = font_name

                    if edu.field_of_study:
                        r_field = p_edu.add_run(f" in {edu.field_of_study}")
                        r_field.font.size = Pt(10)
                        r_field.font.name = font_name

                    r_sch = p_edu.add_run(f" — {edu.school}")
                    r_sch.font.size = Pt(9.5)
                    r_sch.font.name = font_name

                    if edu.end_date:
                        r_dt = p_edu.add_run(f" ({edu.end_date})")
                        r_dt.font.size = Pt(9)
                        r_dt.font.italic = True
                        r_dt.font.name = font_name

                    if edu.gpa:
                        p_gpa = doc.add_paragraph()
                        p_gpa.paragraph_format.space_before = Pt(0)
                        p_gpa.paragraph_format.space_after = Pt(1)
                        r_g = p_gpa.add_run(f"GPA: {edu.gpa}")
                        r_g.font.size = Pt(9)
                        r_g.font.name = font_name

                    if edu.highlights:
                        p_hi = doc.add_paragraph()
                        p_hi.paragraph_format.space_before = Pt(0)
                        p_hi.paragraph_format.space_after = Pt(2)
                        r_h = p_hi.add_run(edu.highlights)
                        r_h.font.size = Pt(9)
                        r_h.font.name = font_name

            elif section_key == "certifications" and data.certifications:
                add_heading("Certifications")
                for cert in data.certifications:
                    p_c = doc.add_paragraph(style='List Bullet')
                    p_c.paragraph_format.space_after = Pt(1)
                    r = p_c.add_run(f"{cert.name} — {cert.issuer} ({cert.date})")
                    r.font.name = font_name
                    r.font.size = Pt(9.5)

            elif section_key == "achievements" and data.achievements:
                add_heading("Achievements & Activities")
                for ach in data.achievements:
                    p_a = doc.add_paragraph(style='List Bullet')
                    p_a.paragraph_format.space_after = Pt(1)
                    r = p_a.add_run(f"{ach.title}: {ach.description}")
                    r.font.name = font_name
                    r.font.size = Pt(9.5)

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream.getvalue()

    def export_pdf(self, data: ResumeBuilderData, template_id: str = "modern") -> bytes:
        """
        Compiles ResumeBuilderData into a clean vector PDF document with selectable text,
        margins, rule dividers, and automatic multi-page overflow handling using PyMuPDF.
        """
        doc = fitz.open()
        
        # Dimensions for Standard US Letter (612 x 792 pt)
        page_width = 612.0
        page_height = 792.0
        margin_x = 46.0
        margin_top = 46.0
        margin_bottom = 46.0
        content_width = page_width - (2 * margin_x)
        
        page = doc.new_page(width=page_width, height=page_height)
        y = margin_top

        # Font configuration
        font_regular = "helv" if template_id != "classic" else "times"
        font_bold = "hebo" if template_id != "classic" else "tibo"

        def check_overflow(needed_space: float = 24.0) -> None:
            nonlocal page, y
            if y + needed_space > page_height - margin_bottom:
                page = doc.new_page(width=page_width, height=page_height)
                y = margin_top

        # 1. Candidate Full Name
        name_text = data.contact.full_name or "Candidate Name"
        name_size = 18.0
        check_overflow(name_size + 10)
        
        if template_id == "classic":
            # Center name
            text_len = fitz.get_text_length(name_text, fontname=font_bold, fontsize=name_size)
            x_pos = max(margin_x, (page_width - text_len) / 2)
            page.insert_text(fitz.Point(x_pos, y + name_size), name_text, fontname=font_bold, fontsize=name_size, color=(0.1, 0.1, 0.2))
        else:
            color = (0.1, 0.2, 0.5) if template_id == "modern" else (0.1, 0.1, 0.1)
            page.insert_text(fitz.Point(margin_x, y + name_size), name_text, fontname=font_bold, fontsize=name_size, color=color)
        y += name_size + 6

        # 2. Contact Information Line
        contact_parts = []
        if data.contact.location: contact_parts.append(data.contact.location)
        if data.contact.phone: contact_parts.append(data.contact.phone)
        if data.contact.email: contact_parts.append(data.contact.email)
        if data.contact.linkedin: contact_parts.append(data.contact.linkedin.replace("https://", ""))
        if data.contact.github: contact_parts.append(data.contact.github.replace("https://", ""))
        
        contact_str = " | ".join(contact_parts)
        if contact_str:
            check_overflow(14)
            c_size = 9.0
            if template_id == "classic":
                c_len = fitz.get_text_length(contact_str, fontname=font_regular, fontsize=c_size)
                x_pos = max(margin_x, (page_width - c_len) / 2)
                page.insert_text(fitz.Point(x_pos, y + c_size), contact_str, fontname=font_regular, fontsize=c_size, color=(0.3, 0.3, 0.3))
            else:
                page.insert_text(fitz.Point(margin_x, y + c_size), contact_str, fontname=font_regular, fontsize=c_size, color=(0.3, 0.3, 0.3))
            y += c_size + 8

        # Header horizontal divider
        page.draw_line(fitz.Point(margin_x, y), fitz.Point(page_width - margin_x, y), color=(0.75, 0.8, 0.85), width=0.75)
        y += 12

        def draw_section_header(title: str):
            nonlocal y
            check_overflow(26)
            h_size = 10.5
            color = (0.1, 0.2, 0.5) if template_id == "modern" else (0.05, 0.05, 0.05)
            page.insert_text(fitz.Point(margin_x, y + h_size), title.upper(), fontname=font_bold, fontsize=h_size, color=color)
            y += h_size + 3
            page.draw_line(fitz.Point(margin_x, y), fitz.Point(page_width - margin_x, y), color=(0.85, 0.85, 0.85), width=0.5)
            y += 8

        # Render sections
        for section_key in data.section_order:
            if section_key == "summary" and data.summary.strip():
                draw_section_header("Professional Summary")
                check_overflow(30)
                rect = fitz.Rect(margin_x, y, margin_x + content_width, y + 80)
                # Textbox wrapping
                inserted_rc = page.insert_textbox(rect, data.summary.strip(), fontname=font_regular, fontsize=9.0, color=(0.15, 0.15, 0.15), lineheight=1.3)
                if inserted_rc < 0:
                    check_overflow(50)
                y += max(35, abs(inserted_rc) if inserted_rc != 0 else 35) + 6

            elif section_key == "skills":
                s = data.skills
                has_skills = any([s.languages, s.frameworks, s.databases, s.tools, s.soft_skills])
                if has_skills:
                    draw_section_header("Technical Skills")
                    skill_rows = []
                    if s.languages: skill_rows.append(("Languages", s.languages))
                    if s.frameworks: skill_rows.append(("Frameworks", s.frameworks))
                    if s.databases: skill_rows.append(("Databases & Cloud", s.databases))
                    if s.tools: skill_rows.append(("Tools", s.tools))
                    
                    for lbl, val in skill_rows:
                        check_overflow(14)
                        line = f"{lbl}: {val}"
                        page.insert_text(fitz.Point(margin_x, y + 9.0), line, fontname=font_regular, fontsize=9.0, color=(0.15, 0.15, 0.15))
                        y += 13
                    y += 4

            elif section_key == "experience" and data.experience:
                draw_section_header("Work Experience")
                for exp in data.experience:
                    check_overflow(22)
                    header_line = f"{exp.role} - {exp.company}"
                    page.insert_text(fitz.Point(margin_x, y + 9.5), header_line, fontname=font_bold, fontsize=9.5, color=(0.1, 0.1, 0.1))
                    
                    if exp.start_date or exp.end_date:
                        date_str = f"{exp.start_date} - {exp.end_date if not exp.is_current else 'Present'}"
                        d_len = fitz.get_text_length(date_str, fontname=font_regular, fontsize=8.5)
                        page.insert_text(fitz.Point(page_width - margin_x - d_len, y + 9.5), date_str, fontname=font_regular, fontsize=8.5, color=(0.4, 0.4, 0.4))
                    y += 14

                    for bullet in exp.bullets:
                        if bullet.strip():
                            check_overflow(18)
                            bullet_str = f"•   {bullet.strip()}"
                            rect = fitz.Rect(margin_x + 8, y, margin_x + content_width, y + 36)
                            page.insert_textbox(rect, bullet_str, fontname=font_regular, fontsize=8.5, color=(0.2, 0.2, 0.2), lineheight=1.25)
                            y += 15
                    y += 4

            elif section_key == "projects" and data.projects:
                draw_section_header("Selected Projects")
                for proj in data.projects:
                    check_overflow(20)
                    p_line = proj.name + (f" ({proj.tech_stack})" if proj.tech_stack else "")
                    page.insert_text(fitz.Point(margin_x, y + 9.5), p_line, fontname=font_bold, fontsize=9.5, color=(0.1, 0.1, 0.1))
                    y += 13

                    for bullet in proj.bullets:
                        if bullet.strip():
                            check_overflow(18)
                            bullet_str = f"•   {bullet.strip()}"
                            rect = fitz.Rect(margin_x + 8, y, margin_x + content_width, y + 36)
                            page.insert_textbox(rect, bullet_str, fontname=font_regular, fontsize=8.5, color=(0.2, 0.2, 0.2), lineheight=1.25)
                            y += 14
                    y += 4

            elif section_key == "education" and data.education:
                draw_section_header("Education")
                for edu in data.education:
                    check_overflow(18)
                    deg_line = f"{edu.degree}" + (f" in {edu.field_of_study}" if edu.field_of_study else "") + f" - {edu.school}"
                    page.insert_text(fitz.Point(margin_x, y + 9.0), deg_line, fontname=font_bold, fontsize=9.0, color=(0.1, 0.1, 0.1))
                    if edu.end_date:
                        d_len = fitz.get_text_length(edu.end_date, fontname=font_regular, fontsize=8.5)
                        page.insert_text(fitz.Point(page_width - margin_x - d_len, y + 9.0), edu.end_date, fontname=font_regular, fontsize=8.5, color=(0.4, 0.4, 0.4))
                    y += 13

                    if edu.highlights or edu.gpa:
                        detail = f"GPA: {edu.gpa} | {edu.highlights}" if edu.gpa else edu.highlights
                        page.insert_text(fitz.Point(margin_x + 8, y + 8.5), detail, fontname=font_regular, fontsize=8.5, color=(0.3, 0.3, 0.3))
                        y += 12
                    y += 3

            elif section_key == "certifications" and data.certifications:
                draw_section_header("Certifications")
                for cert in data.certifications:
                    check_overflow(14)
                    c_line = f"•   {cert.name} - {cert.issuer}" + (f" ({cert.date})" if cert.date else "")
                    page.insert_text(fitz.Point(margin_x + 4, y + 8.5), c_line, fontname=font_regular, fontsize=8.5, color=(0.2, 0.2, 0.2))
                    y += 12
                y += 4

            elif section_key == "achievements" and data.achievements:
                draw_section_header("Achievements")
                for ach in data.achievements:
                    check_overflow(14)
                    a_line = f"•   {ach.title}: {ach.description}"
                    page.insert_text(fitz.Point(margin_x + 4, y + 8.5), a_line, fontname=font_regular, fontsize=8.5, color=(0.2, 0.2, 0.2))
                    y += 12
                y += 4

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

resume_export_service = ResumeExportService()
