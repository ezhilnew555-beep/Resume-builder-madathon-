import io
import pytest
import fitz
import docx
import jwt
from fastapi.testclient import TestClient
from app.main import app
from app.core.config import settings

@pytest.fixture
def client():
    return TestClient(app)

def make_test_token(user_id: str, email: str = "test@example.com") -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "role": "authenticated",
        "aud": "authenticated",
    }
    secret = settings.SUPABASE_JWT_SECRET or "test_secret_for_jwt_validation_32_bytes_long"
    return jwt.encode(payload, secret, algorithm="HS256")

def create_synthetic_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    data = doc.tobytes()
    doc.close()
    return data

def create_scanned_image_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    # Draw a blank rectangle/pixmap with no vector text
    pix = fitz.Pixmap(fitz.csRGB, (0, 0, 100, 100), 1)
    page.insert_image(page.rect, pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return data

def create_synthetic_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Alex Candidate", level=0)
    doc.add_paragraph("alex.candidate@example.com | +1 555-0144 | San Francisco, CA")
    
    doc.add_heading("PROFESSIONAL SUMMARY", level=1)
    doc.add_paragraph("Versatile Full Stack Engineer with 4 years of experience building responsive web apps.")
    
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science, Tech University, Graduated 2023")
    
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph("Python, FastAPI, TypeScript, React, PostgreSQL, Docker, Git")
    
    doc.add_heading("WORK EXPERIENCE", level=1)
    doc.add_paragraph("Full Stack Developer at Cloud Solutions Inc (2022 - Present)")
    doc.add_paragraph("Engineered microservices architecture and increased API throughput by 40%.")

    doc.add_heading("PROJECTS", level=1)
    doc.add_paragraph("CareerLens Platform: AI-powered employability and resume diagnostic tool.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

# -------------------------------------------------------------
# 1. Valid PDF Text Extraction
# -------------------------------------------------------------
def test_valid_pdf_upload_and_extraction(client: TestClient):
    synthetic_text = (
        "Alex Candidate\n"
        "alex.candidate@example.com | 555-0144\n\n"
        "PROFESSIONAL SUMMARY\n"
        "Experienced software engineer with deep proficiency in distributed backend systems and frontend frameworks.\n\n"
        "EDUCATION\n"
        "BS Computer Science, University of Technology, 2024\n\n"
        "TECHNICAL SKILLS\n"
        "Python, FastAPI, React, TypeScript, PostgreSQL, Docker\n\n"
        "WORK EXPERIENCE\n"
        "Software Engineer, Tech Corp (2022 - 2024)\n"
        "Built resilient event-driven data streaming pipelines.\n\n"
        "PROJECTS\n"
        "CareerLens Application\n"
        "Designed an automated resume auditing platform.\n"
    )
    pdf_bytes = create_synthetic_pdf(synthetic_text)
    
    files = {"file": ("alex_resume.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "pdf"
    assert data["parser_method"] == "pymupdf"
    assert data["extraction_status"] == "success"
    assert data["is_scanned"] is False
    assert data["word_count"] > 30
    assert "resume_id" in data
    assert "education" in data["detected_sections"]
    assert "skills" in data["detected_sections"]
    assert "experience" in data["detected_sections"]

    resume_id = data["resume_id"]

    # Verify Detail Endpoint
    detail_res = client.get(f"/api/v1/resume/{resume_id}")
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert "education" in detail["sections"]
    assert "skills" in detail["sections"]
    assert detail["quality_assessment"]["quality_score"] > 50

    # Verify Raw Text Endpoint
    text_res = client.get(f"/api/v1/resume/{resume_id}/text")
    assert text_res.status_code == 200
    text_data = text_res.json()
    assert "Alex Candidate" in text_data["raw_text"]
    assert text_data["word_count"] > 30

# -------------------------------------------------------------
# 2. Valid DOCX Document Extraction
# -------------------------------------------------------------
def test_valid_docx_upload_and_extraction(client: TestClient):
    docx_bytes = create_synthetic_docx()
    files = {"file": ("candidate_profile.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "docx"
    assert data["parser_method"] == "python-docx"
    assert data["extraction_status"] == "success"
    assert "education" in data["detected_sections"]
    assert "skills" in data["detected_sections"]
    assert "experience" in data["detected_sections"]

    resume_id = data["resume_id"]
    text_res = client.get(f"/api/v1/resume/{resume_id}/text")
    assert text_res.status_code == 200
    assert "Cloud Solutions" in text_res.json()["raw_text"]

# -------------------------------------------------------------
# 3. Scanned PDF with Low/Zero Text (OCR Fallback)
# -------------------------------------------------------------
def test_scanned_image_pdf_handles_ocr_fallback(client: TestClient):
    scanned_bytes = create_scanned_image_pdf()
    files = {"file": ("scanned_invoice.pdf", scanned_bytes, "application/pdf")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["is_scanned"] is True
    # If Tesseract is not installed on this host, it must cleanly report ocr_unavailable with setup instructions
    assert data["extraction_status"] in ["ocr_unavailable", "success", "partial"]
    assert "Scanned" in data["status_message"] or "OCR" in data["status_message"]

# -------------------------------------------------------------
# 4. Signature & Format Validation Rejections
# -------------------------------------------------------------
def test_reject_mismatched_file_signature(client: TestClient):
    # Claims to be a PDF, but content is arbitrary plain text
    fake_pdf = b"Hello, this is just arbitrary plain text masquerading as PDF."
    files = {"file": ("fake.pdf", fake_pdf, "application/pdf")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 400
    assert "Invalid file signature" in response.json()["detail"]

def test_reject_unsupported_file_extension(client: TestClient):
    files = {"file": ("resume.txt", b"Plain text resume", "text/plain")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 415
    assert "Unsupported file format" in response.json()["detail"]

def test_reject_macro_enabled_documents(client: TestClient):
    files = {"file": ("resume.docm", b"Macro content", "application/vnd.ms-word.document.macroEnabled.12")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 415
    assert "Macro-enabled" in response.json()["detail"]

def test_reject_empty_file(client: TestClient):
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()

def test_reject_oversized_file(client: TestClient):
    # Create 6 MB byte stream (exceeds default 5 MB limit)
    oversized = b"%PDF-" + b"0" * (6 * 1024 * 1024)
    files = {"file": ("huge.pdf", oversized, "application/pdf")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 413
    assert "exceeds maximum allowed limit" in response.json()["detail"]

def test_reject_corrupted_docx(client: TestClient):
    # Starts with PK but has invalid zip structure
    corrupted_docx = b"PK\x03\x04corrupted_zip_payload"
    files = {"file": ("corrupt.docx", corrupted_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/v1/resume/upload", files=files)

    assert response.status_code == 400
    assert "Corrupted" in response.json()["detail"] or "Malformed" in response.json()["detail"]

# -------------------------------------------------------------
# 5. Access Control & Strict User Isolation
# -------------------------------------------------------------
def test_user_ownership_and_access_control(client: TestClient):
    user_a_id = "11111111-1111-1111-1111-111111111111"
    user_b_id = "22222222-2222-2222-2222-222222222222"

    token_a = make_test_token(user_a_id, "user_a@example.com")
    token_b = make_test_token(user_b_id, "user_b@example.com")

    # User A uploads resume
    pdf_bytes = create_synthetic_pdf("User A Resume Content\n\nEDUCATION\nBS Software\n\nSKILLS\nPython")
    files = {"file": ("user_a_resume.pdf", pdf_bytes, "application/pdf")}
    
    upload_res = client.post(
        "/api/v1/resume/upload",
        files=files,
        headers={"Authorization": f"Bearer {token_a}"}
    )
    assert upload_res.status_code == 200
    resume_id = upload_res.json()["resume_id"]

    # User A can retrieve it
    get_res_a = client.get(
        f"/api/v1/resume/{resume_id}",
        headers={"Authorization": f"Bearer {token_a}"}
    )
    assert get_res_a.status_code == 200

    # User B is blocked with 403 Forbidden
    get_res_b = client.get(
        f"/api/v1/resume/{resume_id}",
        headers={"Authorization": f"Bearer {token_b}"}
    )
    assert get_res_b.status_code == 403
    assert "Access denied" in get_res_b.json()["detail"]

    # Unauthenticated user is also blocked from viewing User A's resume
    get_res_anon = client.get(f"/api/v1/resume/{resume_id}")
    assert get_res_anon.status_code == 403

def test_resume_not_found(client: TestClient):
    response = client.get("/api/v1/resume/00000000-0000-0000-0000-000000000000")
    assert response.status_code == 404
