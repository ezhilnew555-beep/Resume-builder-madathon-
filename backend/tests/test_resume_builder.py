"""
Automated Test Suite for Phase 9: AI Resume Builder, Preview, Export, and Versioning.
Tests draft creation, parsed resume import, autosave, version snapshots,
AI suggestions, native DOCX & PDF exports, re-analysis, and user ownership isolation.
"""

import io
import fitz
from docx import Document
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.core.config import settings
from app.schemas.builder import ResumeBuilderData, ContactData, SkillsData, ProjectEntry, ExperienceEntry, EducationEntry
from tests.test_grounded_explanations import make_test_token

client = TestClient(app)

def test_create_blank_draft():
    res = client.post("/api/v1/builder/resumes", json={"title": "My Software Engineer Resume", "template_id": "modern"})
    assert res.status_code == 200
    data = res.json()
    assert data["draft_id"].startswith("draft-")
    assert data["current_version"] == 1
    assert data["title"] == "My Software Engineer Resume"
    assert data["data"]["template_id"] == "modern"

def test_create_draft_import_parsed_resume():
    # 1. Upload a resume
    doc = Document()
    doc.add_heading("Elena Rostova", 0)
    doc.add_paragraph("Email: elena.rostova@example.com | Phone: 555-123-4567")
    doc.add_paragraph("Summary: Experienced backend software engineer with Python and Go.")
    doc.add_paragraph("Skills: Python, Go, Docker, PostgreSQL, Redis, Kubernetes")
    doc.add_paragraph("Projects: Distributed Cache System built using Go and Redis.")
    doc.add_paragraph("Education: B.S. in Computer Engineering, Tech University, 2023")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    upload_res = client.post(
        "/api/v1/resume/upload",
        files={"file": ("elena_resume.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert upload_res.status_code == 200
    resume_id = upload_res.json()["resume_id"]

    # 2. Create draft importing from parsed resume
    import_payload = {
        "title": "Elena Imported Resume",
        "source_resume_id": resume_id,
        "template_id": "classic"
    }
    draft_res = client.post("/api/v1/builder/resumes", json=import_payload)
    assert draft_res.status_code == 200
    draft_data = draft_res.json()

    assert draft_data["imported_from_resume_id"] == resume_id
    assert "Elena" in draft_data["data"]["contact"]["full_name"]
    assert draft_data["data"]["template_id"] == "classic"
    assert len(draft_data["data"]["projects"]) > 0

def test_update_draft_autosave():
    create_res = client.post("/api/v1/builder/resumes", json={"title": "Draft to Edit"})
    draft_id = create_res.json()["draft_id"]

    # Update summary and skills
    update_payload = {
        "title": "Updated Draft Title",
        "template_id": "minimal",
        "data": {
            "contact": {
                "full_name": "Jordan Smith",
                "email": "jordan@example.com",
                "phone": "555-987-6543",
                "location": "San Francisco, CA"
            },
            "summary": "Results-driven Full Stack Engineer with proven success building resilient APIs.",
            "skills": {
                "languages": "Python, TypeScript",
                "frameworks": "FastAPI, React",
                "databases": "PostgreSQL",
                "tools": "Docker, Git"
            },
            "education": [
                {
                    "id": "edu-1",
                    "school": "UC Berkeley",
                    "degree": "B.S. Computer Science",
                    "end_date": "2024",
                    "gpa": "3.9"
                }
            ],
            "projects": [
                {
                    "id": "proj-1",
                    "name": "Cloud Observability Platform",
                    "tech_stack": "Python, React, Prometheus",
                    "bullets": ["Reduced false positive alerts by 42%."]
                }
            ],
            "experience": [],
            "certifications": [],
            "achievements": [],
            "section_order": ["contact", "summary", "skills", "projects", "education"],
            "template_id": "minimal"
        }
    }

    put_res = client.put(f"/api/v1/builder/resumes/{draft_id}", json=update_payload)
    assert put_res.status_code == 200
    res_data = put_res.json()
    assert res_data["title"] == "Updated Draft Title"
    assert res_data["data"]["contact"]["full_name"] == "Jordan Smith"
    assert res_data["data"]["template_id"] == "minimal"

def test_version_snapshots():
    create_res = client.post("/api/v1/builder/resumes", json={"title": "Versioned Draft"})
    draft_id = create_res.json()["draft_id"]

    # Version 1 is automatically created
    v1_list = client.get(f"/api/v1/builder/resumes/{draft_id}/versions").json()
    assert v1_list["total_versions"] == 1

    # Take snapshot version 2
    snap_res = client.post(
        f"/api/v1/builder/resumes/{draft_id}/version",
        json={"note": "Pre-interview submission polish"}
    )
    assert snap_res.status_code == 200
    assert snap_res.json()["version_number"] == 2

    # Check version list
    v2_list = client.get(f"/api/v1/builder/resumes/{draft_id}/versions").json()
    assert v2_list["total_versions"] == 2
    assert v2_list["versions"][0]["version_number"] == 2
    assert v2_list["versions"][1]["version_number"] == 1

def test_ai_suggestions_clarity_and_metrics():
    # Create draft with weak bullets lacking metrics
    create_payload = {
        "title": "Weak Bullets Draft",
        "initial_data": {
            "contact": {"full_name": "Sam Taylor"},
            "summary": "Worked on backend features and responsible for API development.",
            "skills": {"languages": "Python, SQL", "tools": "Git"},
            "projects": [
                {
                    "id": "proj-w",
                    "name": "Order API",
                    "tech_stack": "Python, FastAPI",
                    "bullets": ["Handled orders and fixed bugs in checkout."]
                }
            ],
            "experience": [
                {
                    "id": "exp-w",
                    "role": "Intern",
                    "company": "Startup Co",
                    "bullets": ["Built endpoints and worked on databases."]
                }
            ],
            "education": [],
            "certifications": [],
            "achievements": [],
            "section_order": ["contact", "summary", "skills", "experience", "projects"],
            "template_id": "modern"
        }
    }
    draft = client.post("/api/v1/builder/resumes", json=create_payload).json()
    draft_id = draft["draft_id"]

    # Request suggestions
    sugg_res = client.post(
        f"/api/v1/builder/resumes/{draft_id}/suggest",
        json={"focus_area": "metrics", "job_description": "We need a Docker engineer."}
    )
    assert sugg_res.status_code == 200
    s_data = sugg_res.json()
    assert len(s_data["suggestions"]) > 0
    # Suggestions must show original vs proposed and explain why
    first_sugg = s_data["suggestions"][0]
    assert len(first_sugg["original_text"]) > 0
    assert len(first_sugg["proposed_text"]) > 0
    assert len(first_sugg["rationale"]) > 0

def test_export_docx_validity():
    create_payload = {
        "title": "Docx Export Test",
        "initial_data": {
            "contact": {
                "full_name": "Marcus Aurelius",
                "email": "marcus@rome.edu",
                "phone": "555-100-2000",
                "location": "Rome, IT"
            },
            "summary": "Philosopher and engineer specializing in distributed governance.",
            "skills": {
                "languages": "Latin, Greek, Python",
                "frameworks": "Stoicism, FastAPI",
                "databases": "PostgreSQL",
                "tools": "Git, Linux"
            },
            "experience": [
                {
                    "id": "exp-1",
                    "role": "Emperor & System Architect",
                    "company": "Roman Administration",
                    "start_date": "161",
                    "end_date": "180",
                    "is_current": False,
                    "bullets": ["Governed distributed empire with zero downtime across 19 provinces."]
                }
            ],
            "projects": [],
            "education": [
                {
                    "id": "edu-1",
                    "school": "Academy of Athens",
                    "degree": "Philosophy & Logic",
                    "end_date": "145",
                    "gpa": "4.0"
                }
            ],
            "certifications": [],
            "achievements": [],
            "section_order": ["contact", "summary", "skills", "experience", "education"],
            "template_id": "classic"
        }
    }
    draft = client.post("/api/v1/builder/resumes", json=create_payload).json()
    draft_id = draft["draft_id"]

    # Export DOCX
    docx_res = client.post(f"/api/v1/builder/resumes/{draft_id}/export/docx")
    assert docx_res.status_code == 200
    assert "wordprocessingml" in docx_res.headers["content-type"]
    assert "attachment" in docx_res.headers["content-disposition"]
    assert len(docx_res.content) > 1000

    # Verify bytes load into python-docx
    doc_opened = Document(io.BytesIO(docx_res.content))
    text_content = "\n".join([p.text for p in doc_opened.paragraphs])
    assert "Marcus Aurelius" in text_content
    assert "Roman Administration" in text_content

def test_export_pdf_validity():
    create_payload = {
        "title": "PDF Export Test",
        "initial_data": {
            "contact": {
                "full_name": "Ada Lovelace",
                "email": "ada@analyticalengine.org",
                "location": "London, UK"
            },
            "summary": "Pioneer of algorithmic programming and computing architectures.",
            "skills": {
                "languages": "Assembly, Mathematical Analysis",
                "tools": "Babbage Analytical Engine"
            },
            "experience": [],
            "projects": [
                {
                    "id": "p-1",
                    "name": "First Computer Algorithm",
                    "tech_stack": "Mathematical Logic",
                    "bullets": ["Computed Bernoulli numbers using mechanical arithmetic loops."]
                }
            ],
            "education": [
                {
                    "id": "edu-1",
                    "school": "University of London",
                    "degree": "Mathematics",
                    "end_date": "1835"
                }
            ],
            "certifications": [],
            "achievements": [],
            "section_order": ["contact", "summary", "skills", "projects", "education"],
            "template_id": "modern"
        }
    }
    draft = client.post("/api/v1/builder/resumes", json=create_payload).json()
    draft_id = draft["draft_id"]

    # Export PDF
    pdf_res = client.post(f"/api/v1/builder/resumes/{draft_id}/export/pdf")
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert len(pdf_res.content) > 1000

    # Verify vector PDF structure via PyMuPDF (fitz)
    pdf_doc = fitz.open(stream=pdf_res.content, filetype="pdf")
    assert pdf_doc.page_count >= 1
    page1_text = pdf_doc[0].get_text()
    assert "Ada Lovelace" in page1_text
    assert "Bernoulli" in page1_text
    pdf_doc.close()

def test_reanalyze_draft_reruns_ats():
    create_payload = {
        "title": "Re-analysis Test",
        "initial_data": {
            "contact": {"full_name": "Robin Vance", "email": "robin@example.com"},
            "summary": "Software engineer with 2 years of experience writing FastAPI and React applications.",
            "skills": {
                "languages": "Python, TypeScript, SQL",
                "frameworks": "FastAPI, React",
                "databases": "PostgreSQL",
                "tools": "Git, Docker"
            },
            "experience": [
                {
                    "id": "exp-1",
                    "role": "Software Engineer",
                    "company": "Tech Corp",
                    "start_date": "2023",
                    "end_date": "2024",
                    "bullets": ["Engineered REST APIs handling 5,000+ daily queries, reducing query response times by 32%."]
                }
            ],
            "projects": [],
            "education": [
                {
                    "id": "edu-1",
                    "school": "State University",
                    "degree": "B.S. in Computer Science",
                    "end_date": "2023",
                    "gpa": "3.85"
                }
            ],
            "certifications": [],
            "achievements": [],
            "section_order": ["contact", "summary", "skills", "experience", "education"],
            "template_id": "modern"
        }
    }
    draft = client.post("/api/v1/builder/resumes", json=create_payload).json()
    draft_id = draft["draft_id"]

    analyze_res = client.post(f"/api/v1/builder/resumes/{draft_id}/analyze")
    assert analyze_res.status_code == 200
    a_data = analyze_res.json()
    assert a_data["draft_id"] == draft_id
    assert 0 <= a_data["overall_score"] <= 100
    assert a_data["detected_skills_count"] >= 3
    assert "content_quality" in a_data["component_scores"]

def test_builder_ownership_isolation():
    token_a = make_test_token("student-alice", "alice@example.com")
    token_b = make_test_token("student-bob", "bob@example.com")

    # Alice creates draft
    alice_res = client.post(
        "/api/v1/builder/resumes",
        json={"title": "Alice Confidential Resume"},
        headers={"Authorization": f"Bearer {token_a}"}
    )
    draft_id = alice_res.json()["draft_id"]

    # Bob attempts to read Alice's draft
    bob_get = client.get(
        f"/api/v1/builder/resumes/{draft_id}",
        headers={"Authorization": f"Bearer {token_b}"}
    )
    assert bob_get.status_code == 403

    # Bob attempts to update Alice's draft
    bob_update = client.put(
        f"/api/v1/builder/resumes/{draft_id}",
        json={"title": "Hacked Title", "data": {"contact": {}, "skills": {}, "section_order": []}},
        headers={"Authorization": f"Bearer {token_b}"}
    )
    assert bob_update.status_code == 403

    # Bob attempts to export Alice's draft
    bob_export = client.post(
        f"/api/v1/builder/resumes/{draft_id}/export/pdf",
        headers={"Authorization": f"Bearer {token_b}"}
    )
    assert bob_export.status_code == 403
